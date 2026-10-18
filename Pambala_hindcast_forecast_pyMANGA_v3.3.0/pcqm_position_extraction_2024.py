"""
2024 PCQM survey (datasets/Millie_German_Pambala_PCQM_2024.xlsx) -> pcqm_2024_normalized.csv -> demo_2024.csv

Step 1 (normalize): read sheet 1 of the workbook (rows from 5; columns A tree type [A adult,
Y young, 0 no tree], B quadrant "<transect>.<point><a-d>", D species, E distance (m), F height (m),
G GBH (cm), H DBH (cm)); DBH = GBH / pi where DBH is missing; species mapped to the five model
species (others kept with species = empty, so vacant and non-model quarters stay in the record).

Step 2 (positions): the same geometrically constrained draw as for 1997 -
    P(r, theta) = delta(r - d_k) * Uniform(alpha_min, alpha_max), epsilon_k ~ N(0, 0.1 m)
for every adult of a model species with distance, DBH and height. Random seed 42.

Output columns of demo_2024.csv: plant, x, y, r_stem (m), h_stem (m), species.

Usage: python pcqm_position_extraction_2024.py
"""
import math
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
XLSX = HERE / "datasets" / "Millie_German_Pambala_PCQM_2024.xlsx"
OUT_NORMALIZED = HERE / "pcqm_2024_normalized.csv"
OUT_DEMO = HERE / "demo_2024.csv"
SEED = 42

PLOT_SPACING = 10.0
TRANSECT_SPACING = 25.0
QUADRANT_ANGLES = {"a": (0, 90), "b": (90, 180), "c": (180, 270), "d": (270, 360)}
DISTANCE_ERROR_STD = 0.1
SPECIES_PREFIX = {"Rhizophora": "Rhizpop", "Avicennia": "Avipop", "Bruguiera": "Brugpop",
                  "Lumnitzera": "Lumnpop", "Excoecaria": "Excopop"}
SPECIES_MAP = {
    "R. muc": "Rhizophora", "R.muc": "Rhizophora", "R. api": "Rhizophora", "R.api": "Rhizophora",
    "L. rac": "Lumnitzera", "L.rac": "Lumnitzera", "L . rac": "Lumnitzera", "L  rac": "Lumnitzera",
    "B. sex": "Bruguiera", "B.sex": "Bruguiera", "B. gym": "Bruguiera", "B.gym": "Bruguiera",
    "E. aga": "Excoecaria", "E.aga": "Excoecaria", "A. off": "Avicennia", "A.off": "Avicennia",
}
HEADER_LABELS = {"Tree", "Species", "Quadrant"}


# ----------------------------------------------------------------------------- step 1
def parse_quadrant(quad_str):
    s = str(quad_str).strip()
    dot = s.index(".")
    return int(s[:dot]), int(s[dot + 1:-1]), s[-1]


def _safe_float(value):
    if value is None:
        return None
    try:
        v = float(value)
        return v if not math.isnan(v) else None
    except (ValueError, TypeError):
        return None


def load_workbook_rows(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for r in range(5, ws.max_row + 1):
        tree_type = ws.cell(row=r, column=1).value
        quadrant = ws.cell(row=r, column=2).value
        if tree_type is None or str(tree_type).strip() in HEADER_LABELS or quadrant is None:
            continue
        q = str(quadrant).strip()
        if q in HEADER_LABELS:
            continue
        species_raw = ws.cell(row=r, column=4).value
        rows.append(dict(tree_type=str(tree_type).strip(), quadrant_raw=q,
                         species_raw=str(species_raw).strip() if species_raw else None,
                         distance_m=_safe_float(ws.cell(row=r, column=5).value),
                         height_m=_safe_float(ws.cell(row=r, column=6).value),
                         gbh_cm=_safe_float(ws.cell(row=r, column=7).value),
                         dbh_cm=_safe_float(ws.cell(row=r, column=8).value)))
    wb.close()
    return rows


def normalize():
    records = []
    for i, row in enumerate(load_workbook_rows(XLSX), start=1):
        try:
            transect, plot, quadrant = parse_quadrant(row["quadrant_raw"])
        except (ValueError, IndexError):
            print("cannot parse quadrant %r (row %d), skipped" % (row["quadrant_raw"], i))
            continue
        species = None if row["species_raw"] is None or row["tree_type"] == "0" else SPECIES_MAP.get(row["species_raw"])
        dbh, gbh = row["dbh_cm"], row["gbh_cm"]
        if species is not None and dbh is None and gbh is not None:
            dbh = gbh / math.pi
        records.append(dict(tree_id=i, tree_type=row["tree_type"], transect=transect, plot=plot, quadrant=quadrant,
                            species_raw=row["species_raw"], species=species, distance_m=row["distance_m"],
                            height_m=row["height_m"], gbh_cm=gbh, dbh_cm=dbh))
    df = pd.DataFrame(records)
    df.to_csv(OUT_NORMALIZED, index=False)
    adults = df[(df["tree_type"] == "A") & df["species"].notna()]
    print("wrote %s: %d records; adults of model species: %s" % (OUT_NORMALIZED.name, len(df), adults["species"].value_counts().to_dict()))
    return df


# ----------------------------------------------------------------------------- step 2
def compute_plot_center(transect, plot):
    return (plot - 1) * PLOT_SPACING, (transect - 1) * TRANSECT_SPACING


def position_tree(x_p, y_p, quadrant, distance, rng):
    alpha_min, alpha_max = QUADRANT_ANGLES[quadrant]
    alpha_rad = math.radians(rng.uniform(alpha_min, alpha_max))
    r = max(0.1, distance + rng.normal(0, DISTANCE_ERROR_STD))
    return x_p + r * math.cos(alpha_rad), y_p + r * math.sin(alpha_rad)


def convert(df, seed=SEED):
    rng = np.random.default_rng(seed)
    adults = df[(df["tree_type"] == "A") & df["species"].notna() & df["distance_m"].notna() & (df["distance_m"] > 0)]
    records = []
    for _, row in adults.iterrows():
        if pd.isna(row["dbh_cm"]) or row["dbh_cm"] <= 0 or pd.isna(row["height_m"]) or row["height_m"] <= 0:
            continue
        x_p, y_p = compute_plot_center(row["transect"], row["plot"])
        x, y = position_tree(x_p, y_p, row["quadrant"], row["distance_m"], rng)
        records.append(dict(x=round(x, 2), y=round(y, 2), r_stem=round(row["dbh_cm"] / 200.0, 4),
                            h_stem=round(row["height_m"], 2), species=row["species"]))
    out = pd.DataFrame(records).sort_values(["species", "x"]).reset_index(drop=True)
    counters, ids = {sp: 0 for sp in SPECIES_PREFIX}, []
    for sp in out["species"]:
        counters[sp] += 1
        ids.append("%s_%09d" % (SPECIES_PREFIX[sp], counters[sp]))
    out.insert(0, "plant", ids)
    out = out[["plant", "x", "y", "r_stem", "h_stem", "species"]]
    out.to_csv(OUT_DEMO, index=False)
    print("wrote %s: %d trees; %s" % (OUT_DEMO.name, len(out), out["species"].value_counts().to_dict()))
    return out


if __name__ == "__main__":
    convert(normalize())
